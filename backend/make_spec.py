"""
심사용 데이터 명세서 생성기.
seeds/*.json 을 읽어 기관별 시트로 된 엑셀 한 부를 만든다.
데이터를 따로 손으로 관리하지 않는다 — 시드가 바뀌면 이 스크립트를 다시 돌린다.

    python make_spec.py            -> ../iM_Worker_Pass_데이터명세서.xlsx
"""

import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SEED_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "seeds")
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..",
                   "iM_Worker_Pass_데이터명세서.xlsx")

SHEETS = [
    ("immigration.json", "출입국", "체류자격·신원"),
    ("bank_kyc.json", "iM뱅크", "급여계좌 실명확인"),
    ("eps.json", "고용센터", "고용이력·사업장 변경"),
    ("clinic.json", "의료기관", "건강진단 유효기간"),
    ("hrdk.json", "산업인력공단", "EPS-TOPIK·취업교육"),
    ("qnet.json", "Q-Net", "국가기술자격"),
    ("kosha.json", "안전보건공단", "안전보건교육"),
    ("niied.json", "교육원", "TOPIK·사회통합프로그램"),
]

LABELS = {
    "worker_name": "성명", "nationality": "국적", "reg_no": "외국인등록번호",
    "birth_date": "생년월일", "gender": "성별", "visa_type": "체류자격",
    "visa_valid_until": "체류 만료일", "account_bank": "은행", "account_number": "계좌번호",
    "passport_no": "여권번호", "kyc_date": "실명확인일",
    "employer_name": "사업장", "employment_from": "근무 시작", "employment_to": "계약 종료",
    "job_category": "직종", "job_change_used": "변경 사용", "job_change_limit": "변경 한도",
    "job_change_excluded": "귀책 미산입", "health_check_date": "건강진단일",
    "health_check_valid_until": "검진 유효기한", "topik_level": "TOPIK 등급",
    "topik_date": "응시일", "visa_status": "체류상태", "industry": "업종",
    "eps_topik_test": "시험", "eps_topik_score": "EPS-TOPIK 점수", "eps_topik_date": "응시일",
    "job_training_name": "취업교육", "job_training_hours": "이수시간", "job_training_date": "이수일",
    "certificate_name": "자격증", "certificate_grade": "등급", "certificate_date": "취득일",
    "issuer": "발급기관", "safety_training_name": "안전보건교육",
    "safety_training_date": "이수일", "safety_training_hours": "이수시간",
    "kiip_program": "사회통합프로그램", "kiip_level": "이수단계",
}

HEAD = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(color="6B7280", size=10)
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load(fn):
    with open(os.path.join(SEED_DIR, fn), encoding="utf-8") as f:
        return json.load(f)


def autosize(ws, rows_start=4):
    widths = {}
    for row in ws.iter_rows(min_row=rows_start):
        for cell in row:
            if cell.value is None:
                continue
            w = sum(2 if ord(ch) > 127 else 1 for ch in str(cell.value)) + 3
            widths[cell.column] = max(widths.get(cell.column, 10), min(w, 46))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def add_sheet(wb, fn, tab, purpose):
    doc = load(fn)
    ws = wb.create_sheet(tab)

    ws["A1"] = f"{doc['_agency']} — {purpose}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"보유 {len(doc['records'])}명 · "
                f"{'실제 발급기관' if doc['_authority'] == 'real' else '모의 API (Phase 2 협력 인터페이스)'} · "
                f"seeds/{fn}")
    ws["A2"].font = NOTE_FONT
    ws["A3"] = doc.get("_note", "")
    ws["A3"].font = NOTE_FONT

    cols = []
    for rec in doc["records"]:
        for k in rec:
            if k not in cols:
                cols.append(k)

    for i, k in enumerate(cols, start=1):
        c = ws.cell(row=5, column=i, value=LABELS.get(k, k))
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center")

    for r, rec in enumerate(doc["records"], start=6):
        for i, k in enumerate(cols, start=1):
            c = ws.cell(row=r, column=i, value=rec.get(k, ""))
            c.border = BORDER
            c.alignment = Alignment(horizontal="center" if k != "employer_name" else "left")

    ws.freeze_panes = "A6"
    autosize(ws, rows_start=5)
    return doc


def add_overview(wb, docs):
    """
    보유 현황 매트릭스 — 행은 사람, 열은 기관. O / X 하나로 읽는다.
    각 기관이 자기 명단만 갖고 있다는 사실이 이 표 하나로 드러난다.
    """
    ws = wb.create_sheet("보유 현황", 0)

    ws["A1"] = "iM Worker Pass — 기관별 데이터 보유 현황"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = ("행은 근로자, 열은 발급기관입니다. O는 그 기관이 해당 인물의 기록을 보유하고 있다는 뜻입니다. "
                "각 기관은 자기 명단만 조회하며, 통합 DB를 나눠 쓰지 않습니다.")
    ws["A2"].font = NOTE_FONT
    ws["A3"] = "모든 인물은 시연용 가상 인물이며 실존 인물과 무관합니다."
    ws["A3"].font = NOTE_FONT

    # 사람 목록은 모든 기관 명단의 합집합. 등장 순서를 유지한다.
    names = []
    for doc, _, _ in docs:
        for rec in doc["records"]:
            if rec["worker_name"] not in names:
                names.append(rec["worker_name"])

    nationality, visa = {}, {}
    for doc, _, _ in docs:
        for rec in doc["records"]:
            if rec.get("nationality"):
                nationality[rec["worker_name"]] = rec["nationality"]
            if rec.get("visa_type"):
                visa[rec["worker_name"]] = (
                    rec["visa_type"] + ("" if rec.get("visa_status") in (None, "유효")
                                        else f" ({rec['visa_status']})"))

    heads = ["성명", "국적", "체류자격"] + [d[0]["_agency"] for d in docs] + ["보유 기관 수"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=5, column=i, value=h)
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 34

    have_fill = PatternFill("solid", fgColor="DCFCE7")
    miss_fill = PatternFill("solid", fgColor="FEF2F2")
    have_font = Font(color="15803D", bold=True)
    miss_font = Font(color="B91C1C")

    for r, name in enumerate(names, start=6):
        ws.cell(row=r, column=1, value=name).border = BORDER
        c = ws.cell(row=r, column=2, value=nationality.get(name, ""))
        c.border, c.alignment = BORDER, Alignment(horizontal="center")
        c = ws.cell(row=r, column=3, value=visa.get(name, ""))
        c.border, c.alignment = BORDER, Alignment(horizontal="center")

        owned = 0
        for i, (doc, _, _) in enumerate(docs, start=4):
            has = name in doc["_index"] if "_index" in doc else any(
                x["worker_name"] == name for x in doc["records"])
            owned += 1 if has else 0
            cell = ws.cell(row=r, column=i, value="O" if has else "X")
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            cell.fill = have_fill if has else miss_fill
            cell.font = have_font if has else miss_font

        c = ws.cell(row=r, column=len(docs) + 4, value=f"{owned} / {len(docs)}")
        c.border, c.alignment = BORDER, Alignment(horizontal="center")

    # 합계 행
    total_row = 6 + len(names)
    c = ws.cell(row=total_row, column=1, value="보유 인원")
    c.fill, c.font, c.border = HEAD, HEAD_FONT, BORDER
    for col in (2, 3):
        ws.cell(row=total_row, column=col, value="").fill = HEAD
        ws.cell(row=total_row, column=col).border = BORDER
    for i, (doc, _, _) in enumerate(docs, start=4):
        c = ws.cell(row=total_row, column=i, value=f"{len(doc['records'])}명")
        c.fill, c.font, c.border = HEAD, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center")
    c = ws.cell(row=total_row, column=len(docs) + 4, value="")
    c.fill, c.border = HEAD, BORDER

    # 권한 구분 행
    auth_row = total_row + 1
    c = ws.cell(row=auth_row, column=1, value="권한 구분")
    c.font, c.border = Font(bold=True, size=10), BORDER
    for col in (2, 3):
        ws.cell(row=auth_row, column=col, value="").border = BORDER
    for i, (doc, _, _) in enumerate(docs, start=4):
        c = ws.cell(row=auth_row, column=i,
                    value="실제 발급기관" if doc["_authority"] == "real" else "모의 API")
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
        c.font = Font(size=9, color="15803D" if doc["_authority"] == "real" else "6B7280")

    ws.freeze_panes = "D6"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    for i in range(4, len(docs) + 4):
        ws.column_dimensions[get_column_letter(i)].width = 15
    ws.column_dimensions[get_column_letter(len(docs) + 4)].width = 13


def main():
    wb = Workbook()
    wb.remove(wb.active)
    docs = []
    for fn, tab, purpose in SHEETS:
        docs.append((add_sheet(wb, fn, tab, purpose), purpose, fn))
    add_overview(wb, docs)
    wb.save(OUT)
    print("생성:", os.path.abspath(OUT))
    for doc, purpose, fn in docs:
        print(f"  {doc['_agency']:20} {len(doc['records'])}명")


if __name__ == "__main__":
    sys.exit(main())
